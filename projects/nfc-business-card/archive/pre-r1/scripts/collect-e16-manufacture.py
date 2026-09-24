#!/usr/bin/env python3
"""Collect the newest E16 manufacturing export out of the EDA download spool.

Every `eda-export-e16-manufacture.js` run drops four files into `~/Downloads`
under random hidden names (`.cn.lceda.pro.XXXXXX`). This script picks them out,
identifies each by content (zip / PDF / two xlsx sheets) and copies them into
`artifacts/manufacture/` with the canonical names that
`check-e16-manufacture.py` expects.

    python3 projects/nfc-business-card/scripts/collect-e16-manufacture.py            # newest four
    python3 projects/nfc-business-card/scripts/collect-e16-manufacture.py --list     # show what it finds
    python3 projects/nfc-business-card/scripts/collect-e16-manufacture.py --spool DIR

Exit status is non-zero when the four roles cannot all be identified, so it is
safe to chain it in front of the checker.
"""

from __future__ import annotations

import argparse
import io
import shutil
import sys
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[3]
TARGET = REPO / "projects/nfc-business-card/artifacts/manufacture"
DEFAULT_SPOOL = Path.home() / "Downloads"
PATTERN = ".cn.lceda.pro.*"

CANONICAL = {
    "gerber": "NFC-E16-gerber.zip",
    "bom": "NFC-E16-bom.xlsx",
    "cpl": "NFC-E16-cpl.xlsx",
    "pdf": "NFC-E16-board-pdf.pdf",
}


def classify(path: Path) -> str | None:
    """Return gerber / bom / cpl / pdf, or None when the file is something else."""
    head = path.read_bytes()[:4]
    if head.startswith(b"%PDF"):
        return "pdf"
    if head[:2] != b"PK":
        return None
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if any(name.endswith(".GTL") for name in names):
                return "gerber"
            if any(name.startswith("xl/") for name in names):
                # The pick-and-place sheet has coordinate columns, the BOM does not.
                # Read through a buffer: the spool files carry random extensions,
                # which openpyxl would otherwise reject.
                import openpyxl

                sheet = openpyxl.load_workbook(io.BytesIO(path.read_bytes()), read_only=True).active
                header = [str(cell.value or "") for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
                for row in sheet.iter_rows(min_row=2, max_row=2):
                    header += [str(cell.value or "") for cell in row]
                joined = " ".join(header).lower()
                if "mid x" in joined or "mid y" in joined or "layer" in joined:
                    return "cpl"
                return "bom"
    except zipfile.BadZipFile:
        return None
    return None


def find_candidates(spool: Path) -> list[Path]:
    return sorted((p for p in spool.glob(PATTERN) if p.is_file()), key=lambda p: p.stat().st_mtime, reverse=True)


def collect(spool: Path, apply: bool) -> int:
    # Exports can run several times in a row and the four files of one run may be
    # written across a couple of minutes (the PDF is slow), so anchor on a gerber
    # and look for one file of each other role inside a +/-5 minute window.
    candidates = sorted(find_candidates(spool), key=lambda p: p.stat().st_mtime)
    kinds = {p: classify(p) for p in candidates}
    gerbers = [p for p in reversed(candidates) if kinds[p] == "gerber"]
    window = 300.0
    chosen: dict[str, Path] = {}
    for anchor in gerbers:
        centre = anchor.stat().st_mtime
        roles: dict[str, Path] = {"gerber": anchor}
        for role in ("bom", "cpl", "pdf"):
            options = [p for p in candidates if kinds[p] == role and abs(p.stat().st_mtime - centre) <= window]
            if options:
                roles[role] = min(options, key=lambda p: abs(p.stat().st_mtime - centre))
        if len(roles) == len(CANONICAL):
            chosen = roles
            break
        print(f"note: gerber {anchor.name} has no complete set (found {sorted(roles)})")
    if not chosen:
        print(f"error: {spool} has no complete export set; re-run eda-export-e16-manufacture.js first", file=sys.stderr)
        return 1
    for role, name in CANONICAL.items():
        source = chosen[role]
        print(f"{role:7s} <- {source.name:28s} {source.stat().st_size / 1024:9.1f} KB")
    if not apply:
        print("dry run: rerun without --list to copy into", TARGET)
        return 0
    TARGET.mkdir(parents=True, exist_ok=True)
    for role, name in CANONICAL.items():
        shutil.copy2(chosen[role], TARGET / name)
    print(f"copied 4 files into {TARGET}")
    print("next: python3 projects/nfc-business-card/scripts/check-e16-manufacture.py")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--spool", type=Path, default=DEFAULT_SPOOL, help="download spool directory")
    parser.add_argument("--list", action="store_true", help="only report what would be copied")
    args = parser.parse_args()
    return collect(args.spool, apply=not args.list)


if __name__ == "__main__":
    raise SystemExit(main())

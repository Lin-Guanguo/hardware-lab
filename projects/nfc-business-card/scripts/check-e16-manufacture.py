#!/usr/bin/env python3
"""Validate the exported E16 manufacturing set against the live PCB snapshot.

The set is produced by scripts/eda-export-e16-manufacture.js and copied into
artifacts/manufacture/ as NFC-E16-gerber.zip, NFC-E16-bom.xlsx,
NFC-E16-cpl.xlsx and NFC-E16-board-pdf.pdf. This checker proves that the fab
data really describes the current board instead of an older revision:

* Gerber zip carries every layer the fabricator needs, the outline layer traces
  the L-shaped board with the USB notch, and the via drill count matches the
  snapshot.
* BOM and pick-and-place carry exactly the snapshot's designators, and every
  placement coordinate matches the snapshot within 0.02 mm.

It writes manifest.json next to the inputs with the file hashes.

    python3 scripts/check-e16-manufacture.py
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[3]
DEFAULT_DIR = REPO / "projects/nfc-business-card/artifacts/manufacture"
DEFAULT_SNAPSHOT = REPO / "projects/nfc-business-card/hardware/e16-right-mid-snapshot.json"
MIL = 39.37007874015748

REQUIRED_LAYERS = {
    "top copper": r"\.GTL$",
    "bottom copper": r"\.GBL$",
    "top solder mask": r"\.GTS$",
    "bottom solder mask": r"\.GBS$",
    "top silkscreen": r"\.GTO$",
    "board outline": r"\.GKO$",
    "top paste": r"\.GTP$",
    "plated drill": r"\.DRL$",
}
# Outline vertices the current board must have (x, y) in mm.
OUTLINE_MUST_HAVE = [
    (33.5, 0.0), (33.5, 15.0), (0.0, 15.0), (0.0, 52.0), (84.0, 52.0),
    (84.0, 20.62), (76.704165, 20.62), (76.704165, 11.38), (84.0, 11.38), (84.0, 0.0),
]
# The USB notch edge used to be drawn at x = 77.5 mm; the J1 footprint's board-edge
# line (76.704165) supersedes it, so the old edge must not come back.
OUTLINE_MUST_BE_ABSENT = [(77.5, 20.62), (77.5, 11.38)]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def gerber_points(gerber_text: str) -> list[tuple[float, float]]:
    match = re.search(r"%FSLAX(\d)(\d)Y(\d)(\d)\*%", gerber_text)
    if not match:
        return []
    ints, decimals = int(match.group(1)), int(match.group(2))
    scale = 10 ** decimals
    points = []
    for line in gerber_text.splitlines():
        hits = re.findall(r"X(-?\d+)Y(-?\d+)", line)
        for xs, ys in hits:
            points.append((int(xs) / scale, int(ys) / scale))
    return points


def snapshot_data(path: Path) -> tuple[dict[str, tuple[float, float]], int]:
    data = json.loads(path.read_text())
    refs = {c["ref"]: (c["x"] / MIL, c["y"] / MIL) for c in data["components"]}
    return refs, len(data["vias"])


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dir", type=Path, default=DEFAULT_DIR)
    parser.add_argument("--snapshot", type=Path, default=DEFAULT_SNAPSHOT)
    args = parser.parse_args()

    directory = args.dir
    files = {
        "gerber": directory / "NFC-E16-gerber.zip",
        "bom": directory / "NFC-E16-bom.xlsx",
        "cpl": directory / "NFC-E16-cpl.xlsx",
        "pdf": directory / "NFC-E16-board-pdf.pdf",
    }
    missing = [name for name, path in files.items() if not path.is_file() or path.stat().st_size == 0]
    if missing:
        raise SystemExit(f"missing export files: {missing}")

    refs, via_count = snapshot_data(args.snapshot)
    report: dict[str, object] = {
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "snapshot": str(args.snapshot.relative_to(REPO)),
        "snapshot_designators": len(refs),
        "snapshot_vias": via_count,
        "files": {name: {"bytes": path.stat().st_size, "sha256": sha256(path)} for name, path in files.items()},
        "problems": [],
    }

    def problem(message: str) -> None:
        report["problems"].append(message)  # type: ignore[union-attr]

    # --- Gerber package -----------------------------------------------------
    with zipfile.ZipFile(files["gerber"]) as archive:
        names = archive.namelist()
        report["gerber_members"] = len(names)
        for label, pattern in REQUIRED_LAYERS.items():
            if not any(re.search(pattern, name) for name in names):
                problem(f"gerber is missing the {label} layer")
        outline_name = next((n for n in names if n.endswith("GKO")), None)
        if outline_name:
            points = gerber_points(archive.read(outline_name).decode("utf-8", "ignore"))
            xs = [p[0] for p in points]
            ys = [p[1] for p in points]
            report["outline_bbox_mm"] = [round(min(xs), 3), round(min(ys), 3), round(max(xs), 3), round(max(ys), 3)]
            if not (abs(min(xs)) < 0.01 and abs(min(ys)) < 0.01 and abs(max(xs) - 84) < 0.01 and abs(max(ys) - 52) < 0.01):
                problem(f"outline bbox is {report['outline_bbox_mm']}, expected 0/0/84/52")
            missing_vertices = [
                vertex for vertex in OUTLINE_MUST_HAVE
                if not any(abs(px - vertex[0]) < 0.001 and abs(py - vertex[1]) < 0.001 for px, py in points)
            ]
            if missing_vertices:
                problem(f"outline is missing vertices {missing_vertices}")
            stale_vertices = [
                vertex for vertex in OUTLINE_MUST_BE_ABSENT
                if any(abs(px - vertex[0]) < 0.01 and abs(py - vertex[1]) < 0.01 for px, py in points)
            ]
            if stale_vertices:
                problem(f"outline still carries superseded vertices {stale_vertices}")
        via_drill = next((n for n in names if "Via" in n and n.endswith(".DRL")), None)
        if via_drill:
            hits = re.findall(r"^X[-\d.]+Y[-\d.]+", archive.read(via_drill).decode("utf-8", "ignore"), flags=re.M)
            report["via_drill_hits"] = len(hits)
            if len(hits) != via_count:
                problem(f"via drill has {len(hits)} hits, snapshot has {via_count} vias")
        else:
            problem("gerber is missing the via drill file")

    # --- BOM ----------------------------------------------------------------
    sheet = openpyxl.load_workbook(files["bom"], data_only=True).active
    header = [str(c.value) for c in sheet[1]]
    designator_column = header.index("Designator") + 1
    bom_refs: set[str] = set()
    rows = 0
    for row in sheet.iter_rows(min_row=2, values_only=False):
        cell = row[designator_column - 1].value
        if not cell:
            continue
        rows += 1
        bom_refs.update(part.strip() for part in str(cell).split(",") if part.strip())
    report["bom_rows"] = rows
    report["bom_designators"] = len(bom_refs)
    if bom_refs != set(refs):
        problem(f"BOM designators differ: only in BOM {sorted(bom_refs - set(refs))}, only in snapshot {sorted(set(refs) - bom_refs)}")

    # --- Pick and place -----------------------------------------------------
    sheet = openpyxl.load_workbook(files["cpl"], data_only=True).active
    header = [str(c.value) for c in sheet[1]]
    columns = {name: header.index(name) for name in ("Designator", "Mid X", "Mid Y")}
    cpl_refs: set[str] = set()
    worst = (0.0, None)
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row[columns["Designator"]]:
            continue
        ref = str(row[columns["Designator"]]).strip()
        cpl_refs.add(ref)
        xy = []
        for key in ("Mid X", "Mid Y"):
            raw = str(row[columns[key]]).strip().lower().removesuffix("mm")
            xy.append(float(raw))
        if ref not in refs:
            continue
        dx = xy[0] - refs[ref][0]
        dy = xy[1] - refs[ref][1]
        distance = (dx * dx + dy * dy) ** 0.5
        if distance > worst[0]:
            worst = (distance, ref)
    report["cpl_designators"] = len(cpl_refs)
    report["worst_placement_delta_mm"] = round(worst[0], 4)
    report["worst_placement_ref"] = worst[1]
    if cpl_refs != set(refs):
        problem(f"CPL designators differ: only in CPL {sorted(cpl_refs - set(refs))}, only in snapshot {sorted(set(refs) - cpl_refs)}")
    if worst[0] > 0.02:
        problem(f"{worst[1]} placement is {worst[0]:.3f} mm away from the snapshot")

    if files["pdf"].stat().st_size < 50_000:
        problem("assembly PDF looks truncated")

    report["ok"] = not report["problems"]
    manifest = directory / "manifest.json"
    manifest.write_text(json.dumps(report, ensure_ascii=False, indent=1) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=1), flush=True)
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
